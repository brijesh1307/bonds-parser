"""Metrics, audit export and template history (items 13, 17, 18)."""

from __future__ import annotations

import csv
import io
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.api.test_templates import MAPPING, SLIP_A, SLIP_B, _files, _post


def test_metrics_require_auth_and_count_parses(client: TestClient, anon: TestClient) -> None:
    assert anon.get("/metrics").status_code == 401
    _post(client, "/api/v1/templates", SLIP_A, MAPPING)
    client.post("/api/v1/parse", files=_files(SLIP_B))
    client.post("/api/v1/parse", files=_files(b"not a pdf"))
    r = client.get("/metrics")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/plain")
    text = r.text
    assert 'bonds_parses_total{market="IN",status="PARSED"} 1.0' in text
    assert 'bonds_template_matches_total{template_id="xyz_securities_contract_note"} 1.0' in text
    assert 'bonds_http_requests_total{method="POST",route="/api/v1/parse",status="415"} 1.0' in text
    assert 'bonds_auth_failures_total{reason="MISSING_HEADER"} 1.0' in text
    assert "bonds_templates_active 1.0" in text
    assert 'route="/api/v1/parse"' in text and "slip.pdf" not in text  # route templates, never file names


def test_audit_export_formats(client: TestClient) -> None:
    client.post("/api/v1/parse", files=_files(SLIP_A))
    client.post("/api/v1/parse", files=_files(SLIP_B))

    r = client.get("/api/v1/audit/export", params={"format": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert 'filename="audit.csv"' in r.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert [row["action"] for row in rows] == ["SLIP_PARSED", "SLIP_PARSED"]  # oldest first
    assert json.loads(rows[0]["details"])["status"] == "NEW_TEMPLATE"
    assert rows[1]["prev_hash"] == rows[0]["row_hash"]  # chain is exported intact

    data = client.get("/api/v1/audit/export", params={"format": "json"}).json()
    assert [d["seq"] for d in data] == [1, 2, 3]  # the CSV export itself was audited as #3
    assert data[2]["action"] == "AUDIT_EXPORTED" and data[2]["details"]["count"] == 2

    wb = load_workbook(io.BytesIO(client.get("/api/v1/audit/export", params={"format": "xlsx"}).content))
    assert wb.sheetnames == ["Audit"] and wb["Audit"].max_row == 5  # header + 4 records

    none = client.get("/api/v1/audit/export", params={"format": "json", "from": "2999-01-01"}).json()
    assert none == []


def test_template_history_shows_changes(client: TestClient) -> None:
    _post(client, "/api/v1/templates", SLIP_A, MAPPING)
    v2 = {
        **MAPPING,
        "template_id": "xyz_securities_contract_note",
        "template_name": None,
        "label_map": {**MAPPING["label_map"], "GST No": "deal_id"},
        "keywords": ["XYZ"],  # type: ignore[dict-item]
        "constants": {"platform": "OTC / NSE RFQ"},
        "note": "second",
    }
    assert _post(client, "/api/v1/templates", SLIP_A, v2).status_code == 201  # type: ignore[attr-defined]

    history = client.get("/api/v1/templates/xyz_securities_contract_note/history").json()
    assert [h["version"] for h in history] == [1, 2]
    assert history[0]["changes"] is None and history[1]["note"] == "second"
    changes = history[1]["changes"]
    assert changes["label_map"]["added"] == {"gst no": "deal_id"}
    assert changes["ignore"] == {"added": [], "removed": ["gst no"]}
    assert changes["constants"]["changed"] == {"platform": {"old": "OTC", "new": "OTC / NSE RFQ"}}
    assert changes["fingerprint.keywords"] == {"added": ["XYZ"], "removed": []}
    assert client.get("/api/v1/templates/nope/history").status_code == 404
