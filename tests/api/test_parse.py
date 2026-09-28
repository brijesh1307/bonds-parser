"""API tests for /health and /api/v1/parse (docs/04_api_spec.md §4.1-4.2, docs/10_testing_strategy.md §6)."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.main import create_app
from tests.conftest import make_pdf

PROBLEM = "application/problem+json"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def client() -> TestClient:
    return TestClient(create_app(Settings.from_env({})))


def _upload(client: TestClient, data: bytes, name: str = "slip.pdf", **params: str) -> object:
    return client.post("/api/v1/parse", files={"file": (name, data, "application/pdf")}, params=params)


def test_health(client: TestClient) -> None:
    r = client.get("/health")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "ok"
    assert body["db"] == "not_configured"
    assert body["version"]


def test_parse_mock01_returns_deal_with_positions(client: TestClient, slip01: Path) -> None:
    data = slip01.read_bytes()
    r = _upload(client, data, name=slip01.name)
    assert r.status_code == 200, r.text  # type: ignore[attr-defined]
    body = r.json()  # type: ignore[attr-defined]

    assert body["status"] == "NEW_TEMPLATE"  # no templates until PH3
    assert body["source"] == {
        "file": slip01.name,
        "pages": 1,
        "page_sizes": [[595.3, 841.9]],
        "sha256": hashlib.sha256(data).hexdigest(),
    }
    deal = body["deal"]
    assert deal["deal_id"] == "GS/NDSOM/2026/004571"
    assert deal["trade_date"] == "2026-09-24"
    assert deal["face_value"] == "50000000"  # decimal as plain string (D10)
    assert deal["price"] == "101.2350"
    assert deal["yield"] == "6.8865"  # alias of Python attribute yield_
    assert deal["consideration"] == "52264305.56"
    assert deal["accrued_days"] == 167
    assert deal["broker"] is None  # nulls are always emitted
    assert "repo" in deal

    fv = body["fields"]["face_value"]
    assert fv["value"] == "50000000"
    assert fv["raw"] == "5,00,00,000.00"
    assert (fv["method"], fv["confidence"], fv["label"], fv["page"]) == (
        "synonym",
        0.95,
        "Face Value (INR)",
        1,
    )
    assert len(fv["bbox"]) == 4

    assert len(body["key_values"]) == 25  # 24 table pairs + the ISIN found in the text (prose)
    mapped = {kv["key"]: kv["mapped_to"] for kv in body["key_values"]}
    assert mapped["Deal Reference No."] == "deal_id"
    assert mapped["Trade Time"] is None
    assert {kv["key"] for kv in body["unmapped"]} == {"Trade Time", "SGL / CSGL A/c"}
    assert body["missing_required"] == []
    assert (deal["deal_type"], deal["buy_sell"], deal["instrument_type"]) == ("OUTRIGHT", "BUY", "GSEC")
    assert (deal["platform"], deal["settlement_mode"], deal["day_count"]) == ("NDS-OM", "DVP-III", "30/360")
    assert body["market"]["market"] == "IN" and body["market"]["profile_available"] is True
    assert body["validation"] == []  # principal, consideration, accrued checks all pass


def test_scanned_pdf_is_unreadable(client: TestClient) -> None:
    body = _upload(client, make_pdf(text=None)).json()  # type: ignore[attr-defined]
    assert body["status"] == "UNREADABLE"
    assert body["key_values"] == []


@pytest.mark.parametrize(
    ("data", "status", "slug"),
    [
        (b"PK\x03\x04 this is a zip", 415, "unsupported-media-type"),
        (make_pdf(encrypt="secret"), 422, "encrypted-pdf"),
        (b"%PDF-1.4\ngarbage", 422, "invalid-pdf"),
        (make_pdf(pages=21), 422, "too-many-pages"),
    ],
)
def test_bad_files_return_problem_json(client: TestClient, data: bytes, status: int, slug: str) -> None:
    r = _upload(client, data)
    assert r.status_code == status  # type: ignore[attr-defined]
    assert r.headers["content-type"] == PROBLEM  # type: ignore[attr-defined]
    body = r.json()  # type: ignore[attr-defined]
    assert body["type"] == f"urn:bonds-parser:problem:{slug}"
    assert body["status"] == status
    assert body["instance"] == "/api/v1/parse"
    assert body["request_id"]


def test_file_too_large() -> None:
    client = TestClient(create_app(Settings.from_env({"BONDS_MAX_UPLOAD_MB": "1"})))
    r = _upload(client, b"%PDF-1.4\n" + b"0" * (1024 * 1024 + 1))
    assert r.status_code == 413  # type: ignore[attr-defined]
    assert r.json()["type"] == "urn:bonds-parser:problem:payload-too-large"  # type: ignore[attr-defined]


def test_missing_file_and_bad_format_are_422_problems(client: TestClient, slip01: Path) -> None:
    r = client.post("/api/v1/parse")
    assert r.status_code == 422
    assert r.json()["type"] == "urn:bonds-parser:problem:request-validation"
    r2 = _upload(client, slip01.read_bytes(), format="pdf")
    assert r2.status_code == 422  # type: ignore[attr-defined]


def test_unknown_route_is_problem_json(client: TestClient) -> None:
    r = client.get("/api/v1/nope")
    assert r.status_code == 404
    assert r.headers["content-type"] == PROBLEM


def test_swagger_documents_the_endpoints(client: TestClient) -> None:
    spec = client.get("/openapi.json").json()
    assert "/health" in spec["paths"]
    parse = spec["paths"]["/api/v1/parse"]["post"]
    assert parse["tags"] == ["Parse"]
    assert {"200", "413", "415", "422"} <= set(parse["responses"])
    assert "yield" in spec["components"]["schemas"]["Deal"]["properties"]


def test_unexpected_error_is_500_problem_without_internals(
    monkeypatch: pytest.MonkeyPatch, slip01: Path
) -> None:
    def boom(*args: object, **kwargs: object) -> None:
        raise RuntimeError("secret internal detail")

    monkeypatch.setattr("app.services.slip_service.run", boom)
    client = TestClient(create_app(Settings.from_env({})), raise_server_exceptions=False)
    r = _upload(client, slip01.read_bytes())
    assert r.status_code == 500  # type: ignore[attr-defined]
    body = r.json()  # type: ignore[attr-defined]
    assert body["type"] == "urn:bonds-parser:problem:internal-error"
    assert "secret" not in r.text  # type: ignore[attr-defined]


def test_xml_download(client: TestClient) -> None:
    from tests.conftest import mock_pdf

    r = _upload(client, mock_pdf("05_client_letter_mld_individual").read_bytes(), format="xml")
    assert r.status_code == 200  # type: ignore[attr-defined]
    assert r.headers["content-type"].startswith("application/xml")  # type: ignore[attr-defined]
    assert 'filename="slip.xml"' in r.headers["content-disposition"]  # type: ignore[attr-defined]
    import xml.etree.ElementTree as ET

    root = ET.fromstring(r.content)  # type: ignore[attr-defined]  # noqa: S314 - our own output
    assert root.tag == "parse_result" and root.get("schema_version") == "1"
    assert root.findtext("deal/counterparty") == "RAHUL DESHPANDE (MOCK)"
    assert root.findtext("deal/consideration") == "1152985.00"
    assert root.findtext("deal/is_market_linked") == "true"
    assert root.find("deal/deal_id").get("null") == "true"  # type: ignore[union-attr]
    assert len(root.findall("key_values/item")) > 20


def test_xlsx_download(client: TestClient, slip01: Path) -> None:
    import io
    from datetime import datetime

    from openpyxl import load_workbook

    r = _upload(client, slip01.read_bytes(), format="xlsx")
    assert r.status_code == 200  # type: ignore[attr-defined]
    assert r.headers["content-type"] == XLSX  # type: ignore[attr-defined]
    assert 'filename="GS_NDSOM_2026_004571.xlsx"' in r.headers["content-disposition"]  # type: ignore[attr-defined]
    wb = load_workbook(io.BytesIO(r.content))  # type: ignore[attr-defined]
    assert wb.sheetnames == ["Deal", "Fields", "Key Values", "Validation"]
    deal = {row[0]: row[1] for row in wb["Deal"].iter_rows(min_row=2, values_only=True)}
    assert deal["consideration"] == 52264305.56  # a real number, not text
    assert isinstance(deal["trade_date"], datetime)  # a real Excel date
    assert deal["deal_type"] == "OUTRIGHT"
