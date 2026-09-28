"""Excel output (docs/05_data_dictionary_and_outputs.md §7).

Sheets: Deal (field | value, repo flattened as repo.*), Fields (confidence and source of each
value), Key Values (everything read from the slip), Validation. Amounts are real numbers and
dates real Excel dates; text that Excel would treat as a formula is neutralised.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.engine.fields import FIELD_SPECS
from app.schemas.slip import ParseResult

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3B63")
DATE_FORMAT = "dd-mmm-yyyy"


def _safe(value: Any) -> Any:
    """Neutralise text Excel would run as a formula (CSV/formula injection)."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _typed(path: str, value: Any) -> Any:
    """Deal values as Excel types: decimals -> numbers, ISO dates -> dates."""
    spec = FIELD_SPECS.get(path)
    if value is None or spec is None:
        return _safe(value)
    if spec.type == "decimal":
        return Decimal(str(value))
    if spec.type == "date" and isinstance(value, str):
        return date.fromisoformat(value)
    return _safe(value)


def _sheet(ws: Worksheet, header: list[str], rows: list[list[Any]]) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, date):
                cell.number_format = DATE_FORMAT
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 70)
    ws.freeze_panes = "A2"


def _deal_rows(result: ParseResult) -> list[list[Any]]:
    deal = result.deal.model_dump(mode="json", by_alias=True)
    rows: list[list[Any]] = [["status", result.status.value], ["file", result.source.file]]
    for key, value in deal.items():
        if key == "repo":
            for k, v in (value or {}).items():
                rows.append([f"repo.{k}", _typed(f"repo.{k}", v)])
        elif key == "identifiers":
            for k, v in value.items():
                rows.append([f"identifiers.{k}", _safe(v)])
        else:
            rows.append([key, _typed(key, value)])
    return rows


def to_xlsx(result: ParseResult) -> bytes:
    wb = Workbook()
    deal_ws = wb.create_sheet("Deal", 0)
    wb.remove(wb.worksheets[1])  # the default empty sheet
    _sheet(deal_ws, ["field", "value"], _deal_rows(result))

    _sheet(
        wb.create_sheet("Fields"),
        ["field", "value", "confidence", "method", "label on slip", "raw text", "page"],
        [
            [path, _typed(path, f.value), f.confidence, f.method.value, _safe(f.label), _safe(f.raw), f.page]
            for path, f in result.fields.items()
        ],
    )
    _sheet(
        wb.create_sheet("Key Values"),
        ["key", "value", "mapped_to", "source", "page"],
        [
            [_safe(kv.key), _safe(kv.value), kv.mapped_to, kv.source.value, kv.page]
            for kv in result.key_values
        ],
    )
    _sheet(
        wb.create_sheet("Validation"),
        ["code", "severity", "message", "fields", "expected", "actual"],
        [
            [v.code, v.severity.value, _safe(v.message), ", ".join(v.fields), v.expected, v.actual]
            for v in result.validation
        ]
        + [
            ["MISSING_REQUIRED", "ERROR", f"required field not found: {p}", p, None, None]
            for p in result.missing_required
        ],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
