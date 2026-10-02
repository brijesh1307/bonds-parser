"""End-to-end demo with the fictitious mock slips - no server, no real data, nothing kept.

Runs the API in-process against a temporary folder and shows:
  1. parsing each mock slip (G-Sec, corporate bond, T-Bill auction, repo, MLD letters)
  2. onboarding the MLD letter layout from one slip, then a different slip in that layout -> PARSED
  3. JSON / XML / Excel outputs
  4. the audit log and its hash-chain verification

Usage:  python tools/demo.py
"""

from __future__ import annotations

import base64
import io
import json
import sys
import tempfile
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.config import Settings  # noqa: E402
from app.main import create_app  # noqa: E402

MOCKS = ROOT / "samples" / "mock"


def _post(client: TestClient, path: str, pdf: Path, **extra: object) -> object:
    with pdf.open("rb") as f:
        return client.post(path, files={"file": (pdf.name, f, "application/pdf")}, **extra)  # type: ignore[arg-type]


def main() -> int:
    with tempfile.TemporaryDirectory() as tmp:
        settings = Settings.from_env(
            {"BONDS_DATA_DIR": f"{tmp}/data", "BONDS_TEMPLATES_DIR": f"{tmp}/templates"}
        )
        app = create_app(settings)
        secret = app.state.client_store.add("demo")
        token = base64.b64encode(f"demo:{secret}".encode()).decode()
        client = TestClient(app, headers={"Authorization": f"Basic {token}", "X-Actor-Name": "Demo User"})

        print("1. Parse every mock slip (no templates yet)\n")
        for pdf in sorted(MOCKS.glob("*.pdf")):
            r = _post(client, "/api/v1/parse", pdf).json()  # type: ignore[attr-defined]
            d = r["deal"]
            print(
                f"   {pdf.stem:38} {r['status']:13} {d['deal_type'] or '-':15} {d['isin']}  "
                f"consideration {d['consideration']}  issues {len(r['validation'])}"
            )

        print("\n2. Onboard the MLD client-letter layout once, from slip 05\n")
        a, b = MOCKS / "05_client_letter_mld_individual.pdf", MOCKS / "06_client_letter_mld_company.pdf"
        mapping = {"template_name": "MLD client confirmation letter", "note": "demo onboarding"}
        prev = _post(client, "/api/v1/templates/preview", a, data={"mapping": json.dumps(mapping)}).json()  # type: ignore[attr-defined]
        print(f"   preview would give: {prev['status']} (missing {prev['missing_required']})")
        saved = _post(client, "/api/v1/templates", a, data={"mapping": json.dumps(mapping)}).json()  # type: ignore[attr-defined]
        tpl = saved["template"]
        print(f"   approved: template {tpl['id']} v{tpl['current_version']} (market {tpl['market']})")
        r = _post(client, "/api/v1/parse", b).json()  # type: ignore[attr-defined]
        d = r["deal"]
        print(
            f"   slip 06 now: {r['status']} via {r['template']['template_id']} "
            f"(score {r['template']['score']})"
        )
        print(
            f"     {d['buy_sell']} {d['quantity']} x {d['face_value_per_unit']} = {d['face_value']}  "
            f"price {d['price']}  consideration {d['consideration']} (stamp duty {d['stamp_duty']})"
        )
        print(f"     counterparty {d['counterparty']}  PAN {d['counterparty_pan']}")

        print("\n3. Outputs\n")
        xml = _post(client, "/api/v1/parse?format=xml", b)
        print(f"   XML : {len(xml.content)} bytes, {xml.headers['content-disposition']}")  # type: ignore[attr-defined]
        xlsx = _post(client, "/api/v1/parse?format=xlsx", b)
        wb = load_workbook(io.BytesIO(xlsx.content))  # type: ignore[attr-defined]
        print(f"   XLSX: sheets {wb.sheetnames}")

        print("\n4. Audit log (metadata only)\n")
        page = client.get("/api/v1/audit", params={"limit": 3}).json()
        for e in page["items"]:
            print(
                f"   #{e['seq']:<3} {e['action']:22} by {e['actor_id']} ({e['actor_name']}) "
                f"status={e['details'].get('status')}"
            )
        print(f"   verify: {client.get('/api/v1/audit/verify').json()}")
        stored = sorted(p.relative_to(tmp).as_posix() for p in Path(tmp).rglob("*") if p.is_file())
        print(f"\n   files on disk after the demo: {stored}  (no PDFs, no deals)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
